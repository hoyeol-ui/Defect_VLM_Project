from pathlib import Path
from docx import Document
from pypdf import PdfReader
import json,csv,statistics,hashlib,re
W=Path(__file__).parent;R=W.parent.parent;S=R/'final_test_evaluation_20260925/statistics'
out=R/'심사용논문_최종test반영본_20260925.docx';d=Document(out)
manifest=json.loads((W/'source_manifest.json').read_text(encoding='utf-8'))
src=Path(manifest['source_docx']);old=Document(src)
checks={};count=0
assert hashlib.sha256(src.read_bytes()).hexdigest()==manifest['source_sha256']
checks['original_docx_unchanged']=True
for file,sha in manifest['ledger_files'].items():assert hashlib.sha256((R/'final_test_evaluation_20260925'/file).read_bytes()).hexdigest()==sha,file
checks['raw_and_statistics_ledgers_unchanged']=True
def rd(n):return list(csv.DictReader((S/n).open(encoding='utf-8-sig')))
def find(rows,**kw):return next(x for x in rows if all(x[k]==str(v) for k,v in kw.items()))
def pc(v):return f'{float(v)*100:.2f}'
def f(v):return f'{float(v):.2f}'
def eq(a,b):
 global count
 assert a==b,(a,b)
 count+=1
body=[p.text for p in d.paragraphs]
for p in old.paragraphs[79:135]:
 if p.text.strip():assert p.text in body,p.text
checks['chapters_1_2_preserved']=True
for i in [151,154,157,158,163,164,166,167,168,170,174,175,177,179,180,182,185,188,189,190,192,194,196,200,201,205,206,207]:
 assert old.paragraphs[i].text in body,i
checks['acquisition_training_and_statistical_design_preserved']=True
for ti in [9,11]:
 for a,b in zip(old.tables[ti].rows,d.tables[ti].rows):
  for ac,bc in zip(a.cells,b.cells):eq(ac.text,bc.text)
checks['selected_composition_tables_preserved']=True
methods=['Random','DINO_F518','Entropy'];datasets=['GC10-DET','NEU-DET']
bs=rd('budget_summary.csv');ss=rd('seed_nAULC.csv');cs=rd('R5_class_AP_by_seed.csv');cd=rd('R5_class_AP_difference.csv');ps=rd('paired_comparisons.csv');fs=rd('R5_vs_full_data.csv');comp=rd('selected_class_composition_summary.csv')
for di,ds in enumerate(datasets):
 for r in range(6):
  for mi,m in enumerate(methods):
   x=find(bs,dataset=ds,method=m,round=r);eq(d.tables[4].cell(1+di*6+r,mi+3).text,f'{pc(x["mean_mAP50_95"])} ± {pc(x["sample_sd_mAP50_95"])}')
 for mi,m in enumerate(methods):
  vals=[float(x['nAULC']) for x in ss if x['dataset']==ds and x['method']==m]
  eq(d.tables[5].cell(di+1,mi+1).text,f'{pc(statistics.mean(vals))} ± {pc(statistics.stdev(vals))}')
  x=find(fs,dataset=ds,method=m);eq(d.tables[7].cell(di+1,mi+2).text,pc(x['r5_mean_mAP50_95']))
 eq(d.tables[7].cell(di+1,1).text,pc(x['full_data_mAP50_95']))
 for oi,other in enumerate(['Random','Entropy']):
  x=find(ps,dataset=ds,method_a='DINO_F518',method_b=other)
  expected=[f(x['paired_mean_difference_pp']),f'{f(x["bootstrap_ci_lower_pp"])} ~ {f(x["bootstrap_ci_upper_pp"])}',x['positive_seeds']+'/5',f'{float(x["exact_two_sided_sign_flip_p"]):.4f}',f'{float(x["holm_adjusted_p"]):.4f}']
  for j,v in enumerate(expected,1):eq(d.tables[6].cell(1+di*2+oi,j).text,v)
 table=d.tables[8 if di==0 else 10];ctable=d.tables[9 if di==0 else 11]
 for i,row in enumerate(table.rows[1:],1):
  c=row.cells[0].text
  for mi,m in enumerate(methods):
   vals=[float(x['R5_AP50_95']) for x in cs if x['dataset']==ds and x['method']==m and x['class_name']==c]
   eq(row.cells[mi+1].text,f'{pc(statistics.mean(vals))} ± {pc(statistics.stdev(vals))}')
   eq(ctable.cell(i,mi+1).text,f'{float(find(comp,dataset=ds,method=m,class_name=c)["mean_selected_images"]):.1f}')
  eq(row.cells[4].text,f(find(cd,dataset=ds,class_name=c)['dino_minus_random_mean_ap_pp']))
checks['numeric_cells_verified']=count
text='\n'.join(body)
for stale in ['성능과 통계는 모두 개발셋','최종평가셋의 평가는 수행하지 않았다','최종평가셋에서의 확인은 이루어지지 않았다','have not yet been confirmed','향후에는 분리해 둔 최종평가셋','26.72%','12.81%p','4.79%p','1.92%p로 다섯 시드 모두에서 음수']:
 assert stale not in text,stale
checks['stale_development_only_claims_absent']=True
assert '다섯 시드 중 네 개에서 음수' in text
assert 'Holm 보정 후 통계적 유의성은 확인되지 않았다' in text
checks['negative_results_and_inference_limits_present']=True
# Tables/figures in the contents must point to the same body pages in final render.
pdf=PdfReader(next((W/'render_delivery').glob('*.pdf')));texts=[p.extract_text() for p in pdf.pages]
mp=json.loads((W/'page_map.json').read_text(encoding='utf-8'))
norm=lambda s:re.sub(r'\s+','',s).replace('−','-').replace('–','-')
for x in mp.values():assert norm(x['title']) in norm(texts[x['pdf_page']-1]),x
checks['contents_targets_verified']=len(mp)
checks['rendered_pages']=len(texts)
checks['output_sha256']=hashlib.sha256(out.read_bytes()).hexdigest()
(W/'final_audit.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(checks,ensure_ascii=False,indent=2))
